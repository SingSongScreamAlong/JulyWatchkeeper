"""
Export Service for WATCHKEEPER

This service handles data export to various formats (PDF, Excel, CSV, JSON).
"""

import os
import csv
import json
from datetime import datetime
from typing import Dict, Any, List
from pathlib import Path
import logging

logger = logging.getLogger(__name__)


class ExportService:
    """Service for exporting data to various formats."""

    def __init__(self):
        """Initialize export service."""
        self.export_dir = Path(os.getenv("EXPORT_DIR", "/tmp/watchkeeper_exports"))
        self.export_dir.mkdir(parents=True, exist_ok=True)

    async def export_intelligence(
        self,
        export_format: str,
        filters: Dict[str, Any]
    ) -> str:
        """
        Export intelligence items.

        Args:
            export_format: Format to export to (pdf, excel, csv, json)
            filters: Filters to apply to the query

        Returns:
            str: Path to the exported file
        """
        from src.core.database import AsyncSessionLocal
        from src.models.intelligence import Intelligence
        from sqlalchemy import select

        async with AsyncSessionLocal() as session:
            # Build query
            query = select(Intelligence)

            # Apply filters (simplified - extend as needed)
            if "status" in filters:
                query = query.where(Intelligence.processing_status == filters["status"])

            # Execute query
            result = await session.execute(query)
            items = result.scalars().all()

            # Convert to list of dicts
            data = []
            for item in items:
                data.append({
                    "id": item.id,
                    "raw_content": item.raw_content[:200] if item.raw_content else "",
                    "processed_content": item.processed_content[:200] if item.processed_content else "",
                    "source_id": item.source_id,
                    "threat_id": item.threat_id,
                    "status": item.processing_status.value if item.processing_status else None,
                    "confidence_score": item.confidence_score,
                    "latitude": item.latitude,
                    "longitude": item.longitude,
                    "created_at": item.created_at.isoformat() if item.created_at else None,
                    "processed_at": item.processed_at.isoformat() if item.processed_at else None,
                })

            # Export based on format
            if export_format == "csv":
                return await self._export_to_csv(data, "intelligence")
            elif export_format == "json":
                return await self._export_to_json(data, "intelligence")
            elif export_format == "excel":
                return await self._export_to_excel(data, "intelligence")
            elif export_format == "pdf":
                return await self._export_to_pdf(data, "intelligence", "Intelligence Items")
            else:
                raise ValueError(f"Unsupported export format: {export_format}")

    async def export_threats(
        self,
        export_format: str,
        filters: Dict[str, Any]
    ) -> str:
        """Export threat items."""
        from src.core.database import AsyncSessionLocal
        from src.models.threat import Threat
        from sqlalchemy import select

        async with AsyncSessionLocal() as session:
            query = select(Threat)

            # Apply filters
            if "category" in filters:
                query = query.where(Threat.category == filters["category"])

            result = await session.execute(query)
            items = result.scalars().all()

            data = []
            for item in items:
                data.append({
                    "id": item.id,
                    "title": item.title,
                    "description": item.description[:200] if item.description else "",
                    "severity": item.severity,
                    "category": item.category.value if item.category else None,
                    "status": item.status.value if item.status else None,
                    "latitude": item.latitude,
                    "longitude": item.longitude,
                    "created_at": item.created_at.isoformat() if item.created_at else None,
                })

            if export_format == "csv":
                return await self._export_to_csv(data, "threats")
            elif export_format == "json":
                return await self._export_to_json(data, "threats")
            elif export_format == "excel":
                return await self._export_to_excel(data, "threats")
            elif export_format == "pdf":
                return await self._export_to_pdf(data, "threats", "Threats")
            else:
                raise ValueError(f"Unsupported export format: {export_format}")

    async def export_sources(
        self,
        export_format: str,
        filters: Dict[str, Any]
    ) -> str:
        """Export source items."""
        from src.core.database import AsyncSessionLocal
        from src.models.source import Source
        from sqlalchemy import select

        async with AsyncSessionLocal() as session:
            query = select(Source)

            # Apply filters
            if "is_active" in filters:
                query = query.where(Source.is_active == filters["is_active"])

            result = await session.execute(query)
            items = result.scalars().all()

            data = []
            for item in items:
                data.append({
                    "id": item.id,
                    "name": item.name,
                    "url": item.url,
                    "source_type": item.source_type.value if item.source_type else None,
                    "is_active": item.is_active,
                    "reliability_score": item.reliability_score,
                    "created_at": item.created_at.isoformat() if item.created_at else None,
                })

            if export_format == "csv":
                return await self._export_to_csv(data, "sources")
            elif export_format == "json":
                return await self._export_to_json(data, "sources")
            elif export_format == "excel":
                return await self._export_to_excel(data, "sources")
            elif export_format == "pdf":
                return await self._export_to_pdf(data, "sources", "Sources")
            else:
                raise ValueError(f"Unsupported export format: {export_format}")

    async def _export_to_csv(self, data: List[Dict[str, Any]], prefix: str) -> str:
        """Export data to CSV format."""
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        filename = f"{prefix}_{timestamp}.csv"
        filepath = self.export_dir / filename

        if not data:
            # Create empty file
            filepath.touch()
            return str(filepath)

        with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=data[0].keys())
            writer.writeheader()
            writer.writerows(data)

        logger.info(f"Exported {len(data)} rows to CSV: {filepath}")
        return str(filepath)

    async def _export_to_json(self, data: List[Dict[str, Any]], prefix: str) -> str:
        """Export data to JSON format."""
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        filename = f"{prefix}_{timestamp}.json"
        filepath = self.export_dir / filename

        with open(filepath, 'w', encoding='utf-8') as jsonfile:
            json.dump(data, jsonfile, indent=2, ensure_ascii=False)

        logger.info(f"Exported {len(data)} items to JSON: {filepath}")
        return str(filepath)

    async def _export_to_excel(self, data: List[Dict[str, Any]], prefix: str) -> str:
        """Export data to Excel format."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        filename = f"{prefix}_{timestamp}.xlsx"
        filepath = self.export_dir / filename

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = prefix.capitalize()

        if not data:
            wb.save(filepath)
            return str(filepath)

        # Write headers
        headers = list(data[0].keys())
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

        # Write data
        for row_idx, item in enumerate(data, start=2):
            for col_idx, header in enumerate(headers, start=1):
                ws.cell(row=row_idx, column=col_idx, value=item.get(header))

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(filepath)
        logger.info(f"Exported {len(data)} rows to Excel: {filepath}")
        return str(filepath)

    async def _export_to_pdf(
        self,
        data: List[Dict[str, Any]],
        prefix: str,
        title: str
    ) -> str:
        """Export data to PDF format."""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet

        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        filename = f"{prefix}_{timestamp}.pdf"
        filepath = self.export_dir / filename

        doc = SimpleDocTemplate(str(filepath), pagesize=A4)
        elements = []

        # Styles
        styles = getSampleStyleSheet()

        # Title
        title_para = Paragraph(f"<b>{title}</b><br/>Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}", styles['Title'])
        elements.append(title_para)
        elements.append(Spacer(1, 20))

        if not data:
            elements.append(Paragraph("No data to display", styles['Normal']))
        else:
            # Prepare table data (limit columns for PDF readability)
            headers = list(data[0].keys())[:6]  # Limit to 6 columns
            table_data = [headers]

            for item in data[:100]:  # Limit to 100 rows
                row = [str(item.get(h, ""))[:50] for h in headers]  # Truncate long values
                table_data.append(row)

            # Create table
            table = Table(table_data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))

            elements.append(table)

            if len(data) > 100:
                elements.append(Spacer(1, 12))
                elements.append(Paragraph(f"<i>Showing first 100 of {len(data)} total rows</i>", styles['Italic']))

        doc.build(elements)
        logger.info(f"Exported {len(data)} items to PDF: {filepath}")
        return str(filepath)
